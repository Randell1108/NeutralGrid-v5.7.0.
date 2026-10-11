from __future__ import annotations

from datetime import datetime, timezone
from datetime import timedelta
from pathlib import Path

import pytest
from openpyxl import load_workbook

import new_bot_data_extractor as extractor
from _bot_data_extractor_core import append_to_excel
from neutralgrid.data.workbook_timestamps import normalize_end_time_utc


@pytest.mark.parametrize('end', [46281.25, '2026/09/16 06/00/00', datetime(2026, 9, 16, 6)])
def test_future_workbook_end_time_is_a_typed_canonical_date(tmp_path: Path, end):
    path = tmp_path/'bots.xlsx'
    columns = ['strategy_id', 'start_time_utc', 'end_time_utc', 'duration_hours']
    extractor.append_row_preserving_schema(path, {'strategy_id':1, 'start_time_utc':'2026-09-16 05:00:00', 'end_time_utc':end, 'duration_hours':1}, columns)
    workbook = load_workbook(path)
    assert workbook.active is not None
    cell = workbook.active['C2']
    assert cell.value == datetime(2026, 9, 16, 6)
    assert cell.number_format == 'yyyy-mm-dd hh:mm:ss'
    assert workbook.active.column_dimensions['C'].number_format == cell.number_format


@pytest.mark.parametrize('bad_end', [None, 'unreadable', 25569.00053565972])
def test_end_falls_back_to_start_plus_recorded_duration(tmp_path: Path, bad_end):
    path = tmp_path/'bots.xlsx'
    extractor.append_row_preserving_schema(path, {'strategy_id':1, 'start_time_utc':datetime(2026, 9, 16, 5, tzinfo=timezone.utc), 'end_time_utc':bad_end, 'duration_hours':1.5}, ['strategy_id','start_time_utc','end_time_utc','duration_hours'])
    workbook = load_workbook(path)
    assert workbook.active is not None
    assert workbook.active['C2'].value == datetime(2026,9,16,6,30)


def test_recorded_zero_profits_are_zero_but_absent_profits_stay_missing():
    assert extractor.derive_profit_factor([0.0,0.0]) == 0.0
    assert extractor.derive_profit_factor([]) is None
    assert extractor.derive_profit_factor([-2.0,0.0]) == 0.0
    assert extractor.derive_profit_factor([3.0,-2.0]) == 1.5
    assert extractor.derive_profit_factor([3.0,0.0]) == 999.99


def test_timezone_is_converted_and_valid_end_beats_rounded_duration():
    end = datetime(2026,9,16,6,0,7,tzinfo=timezone(timedelta(hours=-5)))
    assert normalize_end_time_utc(end, '2026-09-16 10:00:00', 1) == datetime(2026,9,16,11,0,7)


@pytest.mark.parametrize('profit', [float('nan'), float('inf')])
def test_nonfinite_profit_observations_are_not_fabricated_zero(profit):
    assert extractor.derive_profit_factor([profit]) is None


@pytest.mark.parametrize('duration', [None, -1, 'unknown', float('nan')])
def test_invalid_end_without_usable_duration_stays_missing(duration):
    assert normalize_end_time_utc('bad', '2026-09-16 10:00:00', duration) is None


def test_legacy_writer_uses_same_format(tmp_path: Path):
    path = tmp_path/'legacy.xlsx'
    append_to_excel(path, {'strategy_id':1,'start_time_utc':'2026-09-16 05:00:00','end_time_utc':None,'duration_hours':1.5})
    workbook = load_workbook(path)
    assert workbook.active is not None
    assert workbook.active['F2'].value == datetime(2026,9,16,6,30)
    assert workbook.active['F2'].number_format == 'yyyy-mm-dd hh:mm:ss'
