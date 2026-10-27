from __future__ import annotations

from pathlib import Path
from unittest.mock import patch

import numpy as np
import pandas as pd
import pytest

from neutralgrid.models.meta_labeler import ACTIVE_SNAPSHOT_META_FEATURES, _is_label_column
from neutralgrid.training.live_outcome_ingestor import LiveOutcomeIngestor
from neutralgrid.training.unified_training_builder import UnifiedTrainingBuilder


def _workbook(path: Path, *, duplicate: bool = False, conflict: bool = False,
              nonfinite: bool = False) -> None:
    bots = pd.DataFrame([{
        "strategy_id": "123", "symbol": "BTCUSDT",
        "candidate_id": "BTCUSDT_20261008_120000_12345678",
        "start_time_utc": "2026-10-08T13:00:00Z",
        "end_time_utc": "2026-10-08T20:15:00Z", "pnl_pct": 3.5,
        "duration_hours": 6.0,
    }])
    features = {name: float(i + 1) for i, name in enumerate(ACTIVE_SNAPSHOT_META_FEATURES)}
    features.update({
        "strategy_id": "123", "symbol": "BTCUSDT",
        "candidate_id": "BTCUSDT_20261008_120000_12345678",
        "range_prob": 0.7, "trend_prob": 0.2, "persistence_prob": 0.8,
        "hmm_artifact_version": "canonical-hmm",
        "hmm_trained_at_utc": "2026-10-21T00:00:00Z",
        "hmm_feature_semantics_version": "canonical-v1",
        "hmm_feature_source": "verified_replay",
        "feature_cutoff_utc": "2026-10-08T12:00:00Z",
    })
    if conflict:
        features["candidate_id"] = "BTCUSDT_20261008_110000_12345678"
    if nonfinite:
        features["ou_halflife"] = np.inf
    meta = pd.DataFrame([features, features] if duplicate else [features])
    with pd.ExcelWriter(path) as writer:
        bots.to_excel(writer, sheet_name="General", index=False)
        meta.to_excel(writer, sheet_name="Meta Features", index=False)


def _ingest(path: Path) -> pd.DataFrame:
    return LiveOutcomeIngestor(path, path.parent / "links", path.parent / "scans").ingest()


def test_canonical_identity_features_and_lineage_without_ledger(tmp_path: Path) -> None:
    path = tmp_path / "bots.xlsx"
    _workbook(path)
    with patch.object(LiveOutcomeIngestor, "_match_via_forensic", side_effect=AssertionError("must not guess")):
        result = _ingest(path)
    row = result.iloc[0]
    assert row.match_method == "canonical_workbook"
    assert row.hmm_artifact_version == "canonical-hmm"
    assert row.ev_score == float(list(ACTIVE_SNAPSHOT_META_FEATURES).index("ev_score") + 1)
    assert pd.Timestamp(row.t1) == pd.Timestamp("2026-10-08T20:15:00Z")
    assert row.t1_source == "recorded_bot_end"
    assert row.duration_hours == 7.25
    assert row.reported_duration_hours == 6.0
    assert row.training_event_id == "live:123"


@pytest.mark.parametrize("problem", ["duplicate", "conflict"])
def test_canonical_conflicts_never_fall_back_to_old_scanner(tmp_path: Path, problem: str) -> None:
    path = tmp_path / "bots.xlsx"
    _workbook(path, **{problem: True})
    with patch.object(LiveOutcomeIngestor, "_match_via_forensic", side_effect=AssertionError("must not guess")):
        result = _ingest(path)
    assert result.iloc[0].match_method == "canonical_conflict"
    assert result.iloc[0].canonical_rejection_reason


def test_hybrid_accepts_canonical_identity_and_rejects_infinite_features(tmp_path: Path) -> None:
    path = tmp_path / "bots.xlsx"
    builder = UnifiedTrainingBuilder()
    for bad in [False, True]:
        _workbook(path, nonfinite=bad)
        result = builder._load_live_outcome_rows(
            expired_bots_path=str(path), linkage_dir=str(tmp_path / "links"),
            scanner_results_dir=str(tmp_path / "scans"),
        )
        assert len(result) == (0 if bad else 1)


def test_two_executions_of_one_candidate_are_not_deduplicated(tmp_path: Path, monkeypatch) -> None:
    builder = UnifiedTrainingBuilder()
    cid = "BTCUSDT_20261008_120000_12345678"
    backtest = pd.DataFrame([{
        "candidate_id": cid, "symbol": "BTCUSDT", "source_class": "backtest",
        "version_gated": False, "is_authoritative": True,
    }])
    live = pd.DataFrame([{
        "candidate_id": cid, "strategy_id": sid, "training_event_id": f"live:{sid}",
        "symbol": "BTCUSDT", "source": "live", "match_method": "canonical_workbook",
    } for sid in ["123", "124"]])
    monkeypatch.setattr(builder, "_load_backtest_rows", lambda: backtest.copy())
    monkeypatch.setattr(builder, "_load_live_outcome_rows", lambda **kwargs: live.copy())
    result = builder.build_meta_labeler_pool(include_live_outcomes=True, max_rows_per_symbol=0)
    assert set(result.strategy_id.dropna().astype(str)) == {"123", "124"}
    assert len(result) == 2  # Simulation for the candidate is replaced by real executions.


def test_reported_duration_is_blocked_as_a_label_column() -> None:
    assert _is_label_column("reported_duration_hours")


def test_canonical_values_replace_stale_direct_snapshot(tmp_path: Path) -> None:
    path = tmp_path / "bots.xlsx"
    _workbook(path)
    links = tmp_path / "links"
    links.mkdir()
    pd.DataFrame([{"strategy_id": "123", "candidate_id": "BTCUSDT_20261008_120000_12345678"}]).to_csv(links / "deploy_linkage_log.csv", index=False)
    with patch.object(LiveOutcomeIngestor, "_extract_scan_features", side_effect=AssertionError("must not read stale snapshot")):
        result = _ingest(path)
    assert result.iloc[0].match_method == "linkage"
    assert result.iloc[0].hmm_artifact_version == "canonical-hmm"


def test_ledger_disagreement_quarantines_canonical_row(tmp_path: Path) -> None:
    path = tmp_path / "bots.xlsx"
    _workbook(path)
    links = tmp_path / "links"
    links.mkdir()
    pd.DataFrame([{"strategy_id": "123", "candidate_id": "BTCUSDT_20261008_110000_12345678"}]).to_csv(links / "deploy_linkage_log.csv", index=False)
    result = _ingest(path)
    assert result.iloc[0].match_method == "canonical_conflict"
    assert result.iloc[0].canonical_rejection_reason == "canonical_ledger_candidate_disagreement"


@pytest.mark.parametrize("cutoff", ["2026-10-08T13:00:01Z", "invalid", 46303.5])
def test_invalid_or_future_cutoff_never_falls_back(tmp_path: Path, cutoff) -> None:
    from openpyxl import load_workbook
    path = tmp_path / "bots.xlsx"
    _workbook(path)
    book = load_workbook(path)
    sheet = book["Meta Features"]
    headers = [cell.value for cell in sheet[1]]
    sheet.cell(2, headers.index("feature_cutoff_utc") + 1, cutoff)
    book.save(path)
    book.close()
    with patch.object(LiveOutcomeIngestor, "_match_via_forensic", side_effect=AssertionError("must not guess")):
        result = _ingest(path)
    assert result.iloc[0].canonical_rejection_reason == "canonical_feature_cutoff_invalid_or_future"


def test_workbook_change_between_general_and_meta_fails_closed(tmp_path: Path, monkeypatch) -> None:
    path = tmp_path / "bots.xlsx"
    _workbook(path)
    original = LiveOutcomeIngestor._load_bots

    def mutate_after_general(self):
        result = original(self)
        _workbook(path, conflict=True)
        return result

    monkeypatch.setattr(LiveOutcomeIngestor, "_load_bots", mutate_after_general)
    with pytest.raises(ValueError, match="changed during identity/feature read"):
        _ingest(path)
