"""Publish only the reviewed 57 feature rows with exact preservation checks."""
from __future__ import annotations
import hashlib
import json
from pathlib import Path
import numpy as np
import openpyxl
import pandas as pd
from pandas.testing import assert_frame_equal
from neutralgrid.models.meta_labeler import ACTIVE_SNAPSHOT_META_FEATURES
from scripts.backfill_training_features import BASELINE_BACKFILL_COLUMNS, LIVE_PLUS_BACKFILL_COLUMNS_V20260312, HMM_LINEAGE_COLUMNS

OUT = Path(__file__).resolve().parent
ROOT = OUT.parents[1]

def main():
    audit = json.loads((OUT/'full_57_completion_audit.json').read_text(encoding='utf-8'))
    assert audit['status'] == 'PASS' and audit['finite_feature_cells'] == 1140
    source = ROOT/'data/new_expired_bots.xlsx'
    backup = OUT/'backups/workbook_before_full_57.xlsx'
    assert source.read_bytes() == backup.read_bytes(), 'Concurrent workbook edit'
    frame = pd.read_excel(OUT/'full_57_completed.xlsx').set_index('strategy_id')
    wanted = set(int(x) for x in json.loads((OUT/'workbook_intake_audit.json').read_text())['new_strategy_ids'])
    assert set(frame.index) == wanted and len(wanted) == 57
    columns = sorted(set(BASELINE_BACKFILL_COLUMNS + LIVE_PLUS_BACKFILL_COLUMNS_V20260312 + HMM_LINEAGE_COLUMNS + list(ACTIVE_SNAPSHOT_META_FEATURES) + ['survival_prob','backfill_status']) & set(frame.columns))
    assert not set(columns) & {'strategy_id','pnl','pnl_pct','start_time_utc','end_time_utc','candidate_id','symbol'}
    frame['backfill_status'] = 'complete'
    before = pd.read_excel(source, sheet_name=None)
    wb = openpyxl.load_workbook(source)
    ws = wb['Meta Features']
    headers = {c.value:c.column for c in ws[1] if c.value is not None}
    for field in columns:
        if field not in headers:
            headers[field] = ws.max_column+1
            ws.cell(1, headers[field], field)
    updated = 0
    for n in range(2, ws.max_row+1):
        ident = int(ws.cell(n, headers['strategy_id']).value)
        if ident not in wanted:
            continue
        for field in columns:
            value = frame.at[ident, field]
            if pd.isna(value):
                value = None
            elif isinstance(value, np.generic):
                value = value.item()
            ws.cell(n, headers[field]).value = value
        updated += 1
    assert updated == 57
    candidate = OUT/'canonical_full_57_candidate.xlsx'
    wb.save(candidate)
    after = pd.read_excel(candidate, sheet_name=None)
    assert list(before) == list(after)
    for sheet, old in before.items():
        if sheet != 'Meta Features':
            assert_frame_equal(old, after[sheet], check_exact=True)
        else:
            mask = old.strategy_id.isin(wanted)
            assert_frame_equal(old.loc[~mask], after[sheet].loc[~mask, old.columns], check_dtype=False, check_exact=True)
            untouched = [c for c in old if c not in columns]
            assert_frame_equal(old[untouched], after[sheet][untouched], check_dtype=False, check_exact=True)
            matrix = after[sheet].loc[mask,list(ACTIVE_SNAPSHOT_META_FEATURES)].to_numpy(dtype=float)
            assert np.isfinite(matrix).all()
    flat = pd.read_excel(OUT/'utility_backfilled_entry_time.xlsx').set_index('strategy_id')
    for field in columns:
        if field not in flat:
            flat[field] = None
        flat.loc[frame.index, field] = frame[field]
    flat.reset_index().to_excel(OUT/'utility_complete_features.xlsx', index=False)
    candidate.replace(source)
    result = {'status':'PASS','published_at_utc':pd.Timestamp.now(tz='UTC').isoformat(),'updated_meta_rows':updated,'total_meta_rows':len(after['Meta Features']),'selected_feature_count':20,'finite_selected_cells':1140,'unchanged_existing_meta_rows':366,'all_other_sheets_and_unrelated_columns_preserved':True,'changed_columns':columns,'before_sha256':hashlib.sha256(backup.read_bytes()).hexdigest(),'after_sha256':hashlib.sha256(source.read_bytes()).hexdigest()}
    (OUT/'full_57_publication_audit.json').write_text(json.dumps(result,indent=2),encoding='utf-8')
    print(json.dumps(result,indent=2))

if __name__ == '__main__':
    main()
